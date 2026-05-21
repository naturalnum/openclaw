#!/usr/bin/env python3
"""Fill the reusable PPT template with content from the Shanghai task book."""

from __future__ import annotations

import argparse
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Pt


TEXT = RGBColor(31, 45, 61)
MUTED = RGBColor(73, 95, 120)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Fill project-report PPT template from task-book content.")
    parser.add_argument("--input", required=True, help="Template PPTX path.")
    parser.add_argument("--output", required=True, help="Output PPTX path.")
    return parser.parse_args()


def set_shape_text(shape, text: str, size: float = 18, bold: bool = False,
                   color: RGBColor = TEXT, align: PP_ALIGN = PP_ALIGN.LEFT) -> None:
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraphs = text.split("\n")
    for idx, item in enumerate(paragraphs):
        p = frame.paragraphs[0] if idx == 0 else frame.add_paragraph()
        p.alignment = align
        run = p.add_run()
        run.text = item
        run.font.name = "Microsoft YaHei"
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
        if idx > 0:
            p.space_before = Pt(2)


def fill_table(shape, values: list[list[str]], size: float = 12) -> None:
    table = shape.table
    for r, row in enumerate(values):
        for c, value in enumerate(row):
            if r >= len(table.rows) or c >= len(table.columns):
                continue
            cell = table.cell(r, c)
            cell.text = value
            for paragraph in cell.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.name = "Microsoft YaHei"
                    run.font.size = Pt(size)


def main() -> None:
    args = parse_args()
    input_path = Path(args.input).expanduser().resolve()
    output_path = Path(args.output).expanduser().resolve()
    output_path.parent.mkdir(parents=True, exist_ok=True)

    prs = Presentation(str(input_path))

    # Slide 1 cover
    slide = prs.slides[0]
    set_shape_text(slide.shapes[0], "基于人机协同的配网检修运行方式智能校核与编制优化关键技术研究", size=24, bold=True)
    set_shape_text(slide.shapes[2], "2025 年 12 月", size=18, color=MUTED, align=PP_ALIGN.CENTER)

    # Slide 3 background
    slide = prs.slides[2]
    set_shape_text(slide.shapes[2], "面向上海超大城市配电网，检修计划编排与运行方式调整面临设备密集、约束繁杂、多专业协同难等现实挑战。", size=18)
    set_shape_text(slide.shapes[9], "传统人工提报、人工校核、人工协调方式效率低、可追溯性弱，难以支撑多时间尺度检修计划与方式预案的高质量协同编制。", size=16)
    set_shape_text(slide.shapes[11], "[插入上海配网检修业务场景图 / 现状问题示意图]", size=16, bold=True, color=MUTED, align=PP_ALIGN.CENTER)

    # Slide 4 objective
    slide = prs.slides[3]
    set_shape_text(slide.shapes[2], "项目聚焦人机协同的智能校核和编制优化关键技术，融合配电网状态计算、安全校核、检修计划报送与预案编制领域知识，建立数据知识融合的智能决策推理与演化模型。", size=17)
    set_shape_text(slide.shapes[8], "形成可信、安全、高效的检修计划与运行方式智能编制能力", size=18)

    # Slide 5 value
    slide = prs.slides[4]
    set_shape_text(slide.shapes[23], "关键痛点：数据复杂、约束缺失、过程低效", size=16, bold=True)
    set_shape_text(slide.shapes[25], "能力方向：先验知识建模与人机数据知识增强", size=14)
    set_shape_text(slide.shapes[28], "风险管控：主动发现、风险预测与溯因分析", size=14)
    set_shape_text(slide.shapes[29], "优化目标：检修计划排期与运行方式编排协同优化", size=14)
    set_shape_text(slide.shapes[31], "交付结果：智能校核与编制优化模块原型", size=14)
    set_shape_text(slide.shapes[33], "预期成效：效率从日级提升至小时级", size=14)

    # Slide 6 project overview tables
    slide = prs.slides[5]
    fill_table(
        slide.shapes[4],
        [
            ["项目名称", "承担单位", "总经费", "实施期限"],
            ["基于人机协同的配网检修运行方式智能校核与编制优化关键技术研究", "国网上海市电力公司", "525 万元", "2024.1-2025.12"],
            ["项目负责人", "主要出资单位", "承担单位数", "研究人员"],
            ["金敏杰", "国网上海市电力公司", "5 家", "131 人"],
        ],
        size=11,
    )
    fill_table(
        slide.shapes[5],
        [
            ["课题", "研究单位"],
            ["课题1：人机协同数据知识增强架构研究", "中国电力科学研究院有限公司"],
            ["课题2：风险主动发现与溯因技术研究", "西安交通大学 / 中国电力科学研究院有限公司"],
            ["课题3：专家反馈趋优学习与知识更新", "北京交通大学 / 中国电力科学研究院有限公司"],
            ["课题4：智能校核与编制优化模块开发应用", "国网上海市电力公司 / 中国电力科学研究院有限公司 / 国网信息通信产业集团有限公司"],
        ],
        size=10.5,
    )

    # Slide 7 org
    slide = prs.slides[6]
    set_shape_text(slide.shapes[2], "牵头单位：国网上海市电力公司\n项目负责人：金敏杰\n课题负责人：乔骥、黄玉雄 / 陈予尧、王小君 / 李家腾、苏运 / 张海涛", size=16)
    set_shape_text(slide.shapes[6], "组织模式：上海公司牵头，课题单位协同，结合专家咨询机制推进实施方案、课题攻关、模块研发与试点验证。", size=15)
    set_shape_text(slide.shapes[10], "[组织架构图占位]", size=14, bold=True, color=MUTED, align=PP_ALIGN.CENTER)
    set_shape_text(slide.shapes[11], "[课题协同关系图占位]", size=14, bold=True, color=MUTED, align=PP_ALIGN.CENTER)
    set_shape_text(slide.shapes[12], "[专家咨询机制占位]", size=14, bold=True, color=MUTED, align=PP_ALIGN.CENTER)

    # Slide 8 management
    slide = prs.slides[7]
    set_shape_text(slide.shapes[1], "项目过程管理", size=18, bold=True)
    set_shape_text(slide.shapes[5], "项目已形成课题实施方案和协同工作机制，围绕实施方案、阶段检查、月度交底、季度讨论和专家咨询评审，动态优化技术路线、任务分工、接口协同和应用验证安排。", size=15)
    set_shape_text(slide.shapes[6], "管理抓手：实施方案 + 阶段检查 + 专家咨询 + 应用验证", size=14)
    for idx in [8, 9, 10, 11, 12, 13]:
        set_shape_text(slide.shapes[idx], "[过程管理占位]", size=12, bold=True, color=MUTED, align=PP_ALIGN.CENTER)

    # Slide 9 timeline
    slide = prs.slides[8]
    set_shape_text(slide.shapes[2], "项目进度安排", size=18, bold=True)
    set_shape_text(slide.shapes[3], "2024.1-2025.12", size=16, bold=True, align=PP_ALIGN.CENTER)
    timeline_labels = {
        16: "2024.1-2024.3",
        18: "2024.4-2024.12",
        19: "2025.1-2025.6",
        20: "2025.7-2025.9",
        21: "2025.10-2025.12",
        22: "形成实施方案",
        24: "核心方法研究与原型开发",
        25: "系统集成与训练测试",
        26: "试点应用验证",
    }
    for idx, value in timeline_labels.items():
        set_shape_text(slide.shapes[idx], value, size=11, align=PP_ALIGN.CENTER)
    set_shape_text(slide.shapes[29], "[项目时间轴 / 关键节点图占位]", size=14, bold=True, color=MUTED, align=PP_ALIGN.CENTER)
    set_shape_text(slide.shapes[30], "[阶段任务分布图占位]", size=14, bold=True, color=MUTED, align=PP_ALIGN.CENTER)

    # Slide 10 changes
    slide = prs.slides[9]
    set_shape_text(slide.shapes[2], "任务书下达稿阶段暂无重大变更事项。后续如发生负责人、经费或实施计划调整，将按国家电网科技项目管理流程履行备案。", size=16)
    for idx in [6, 7, 8, 9, 10]:
        set_shape_text(slide.shapes[idx], "[变更事项预留占位]", size=13, bold=True, color=MUTED, align=PP_ALIGN.CENTER)

    # Slide 11 budget
    slide = prs.slides[10]
    set_shape_text(slide.shapes[2], "项目总预算 525 万元。根据任务书，承担单位预算分配为上海公司 5 万、中国电力科学研究院有限公司 360 万、北京交通大学 60 万、西安交通大学 60 万、国网信息通信产业集团有限公司 40 万。", size=15)
    fill_table(
        slide.shapes[6],
        [
            ["单位", "总预算", "累计支出", "预算执行率"],
            ["国网上海市电力公司", "5", "待更新", "待更新"],
            ["中国电力科学研究院有限公司", "360", "待更新", "待更新"],
            ["国网信息通信产业集团有限公司", "40", "待更新", "待更新"],
            ["北京交通大学", "60", "待更新", "待更新"],
            ["西安交通大学", "60", "待更新", "待更新"],
            ["总计", "525", "待更新", "待更新"],
        ],
        size=11,
    )

    # Slide 17 overall research content
    slide = prs.slides[16]
    set_shape_text(slide.shapes[3], "项目围绕“数据知识增强架构、风险主动发现与溯因、专家反馈趋优学习、智能校核与编制优化模块开发应用”四个方向展开，形成从知识建模到应用验证的闭环研究链路。", size=16)
    set_shape_text(slide.shapes[6], "[整体研究框架图 / 四课题关系图占位]", size=16, bold=True, color=MUTED, align=PP_ALIGN.CENTER)

    # Slide 18 task 1
    slide = prs.slides[17]
    set_shape_text(slide.shapes[7], "课题 1", size=20, bold=True, align=PP_ALIGN.CENTER)
    set_shape_text(slide.shapes[9], "面向配网检修运行方式动态协调编排的人机协同数据知识增强架构研究。研究多时间尺度检修计划编排和运行方式调整的耦合机制、先验知识提取与表征方法，形成人机交互与协同的整体架构。", size=15)
    set_shape_text(slide.shapes[11], "研究内容与目标", size=18, bold=True)
    for idx, value in {
        45: "联合建模方法",
        62: "先验知识提取与表征",
        64: "数据知识融合增强架构",
        66: "考核：专利 2 项、论文 2 篇、技术报告 1 份",
        68: "预期：提出基于人机混合增强智能的配网检修运行方式编排架构",
    }.items():
        set_shape_text(slide.shapes[idx], value, size=11)

    # Slide 19 task 2
    slide = prs.slides[18]
    set_shape_text(slide.shapes[0], "课题 2", size=20, bold=True, align=PP_ALIGN.CENTER)
    set_shape_text(slide.shapes[3], "基于先验知识内嵌机器学习的配网检修风险主动发现与溯因技术研究。围绕安全约束建模、多粒度时空学习、分析-校验双向反馈和风险溯因展开研究。", size=15)
    set_shape_text(slide.shapes[5], "风险主动发现与溯因", size=18, bold=True)
    set_shape_text(slide.shapes[8], "安全约束建模", size=12)
    set_shape_text(slide.shapes[9], "多粒度时空学习", size=12)
    set_shape_text(slide.shapes[10], "分析-校验双向反馈", size=12)
    for idx, label in [(11, "[风险预测技术路线图占位]"), (12, "[风险溯因流程图占位]"), (13, "[关键指标示意图占位]")]:
        set_shape_text(slide.shapes[idx], label, size=14, bold=True, color=MUTED, align=PP_ALIGN.CENTER)

    # Slide 20 task 3
    slide = prs.slides[19]
    set_shape_text(slide.shapes[0], "课题 3", size=20, bold=True, align=PP_ALIGN.CENTER)
    set_shape_text(slide.shapes[2], "检修计划与运行方式预案的专家反馈趋优学习与知识更新技术研究。重点研究智能体构建、偏好强化学习、人机协同双向交互以及知识库构建、推理与更新。", size=15)
    set_shape_text(slide.shapes[4], "专家反馈趋优学习与知识更新", size=18, bold=True)
    set_shape_text(slide.shapes[7], "关键任务：计划编排智能体、运行方式预案生成、调度员评价反馈、知识库构建与更新。", size=13)
    set_shape_text(slide.shapes[23], "考核指标：知识节点数 >10000，安全校核通过率 >95%，专利 3 项，论文 2 篇，技术报告 1 份。", size=12)
    set_shape_text(slide.shapes[26], "阶段重点：2024 年开展规则表示和偏好强化学习研究，2025 年开展人类反馈趋优训练与知识库推理更新。", size=12)
    for idx, label in [(27, "[趋优训练流程图占位]"), (28, "[评价反馈机制图占位]"), (29, "[知识库结构图占位]"), (30, "[更新机制图占位]")]:
        set_shape_text(slide.shapes[idx], label, size=13, bold=True, color=MUTED, align=PP_ALIGN.CENTER)

    # Slide 21 task 4
    slide = prs.slides[20]
    set_shape_text(slide.shapes[0], "课题 4", size=20, bold=True, align=PP_ALIGN.CENTER)
    set_shape_text(slide.shapes[2], "人机协同数据知识增强的配电系统检修运行方式智能校核与编制优化模块的开发与应用。围绕模块原型、基础架构与集成方案、地区电网试点应用验证开展研发。", size=15)
    set_shape_text(slide.shapes[4], "模块开发与应用验证", size=18, bold=True)
    set_shape_text(slide.shapes[7], "[模块架构图 / 系统集成图 / 试点应用场景图占位]", size=15, bold=True, color=MUTED, align=PP_ALIGN.CENTER)

    # Slide 22 overall indicators
    slide = prs.slides[21]
    set_shape_text(slide.shapes[0], "项目指标与交付", size=20, bold=True, align=PP_ALIGN.CENTER)
    set_shape_text(slide.shapes[2], "项目最终交付包括总报告和 4 份技术报告、11 项发明专利、6 篇论文、1 项软件著作权，以及与 PMS 等系统集成的智能校核与编制优化模块试点应用。", size=15)
    set_shape_text(slide.shapes[4], "关键指标", size=18, bold=True)
    set_shape_text(slide.shapes[9], "风险研判类别不少于 3 类；知识节点数 >10000；安全校核通过率 >95%；风险辨识准确率 >90%。", size=12)
    set_shape_text(slide.shapes[10], "整体编制效率由日级提升至小时级，并在 1 个地区电网开展试点应用。", size=12)
    set_shape_text(slide.shapes[12], "[指标卡片 / KPI 图表占位]", size=14, bold=True, color=MUTED, align=PP_ALIGN.CENTER)
    set_shape_text(slide.shapes[13], "[交付成果矩阵图占位]", size=14, bold=True, color=MUTED, align=PP_ALIGN.CENTER)

    # Slide 58 outlook
    slide = prs.slides[57]
    set_shape_text(slide.shapes[1], "项目将围绕任务书既定目标持续推进。下一阶段将重点完成检修计划与方式预案协同编制关键方法研究、模块原型研发、系统集成与地区电网试点验证，并在结题阶段形成专利、论文、软件著作权和完整技术报告成果。", size=16)

    # Slide 60 expert response placeholder
    slide = prs.slides[59]
    set_shape_text(slide.shapes[1], "任务书下达稿阶段暂未形成中期督导或验收专家意见，本页预留用于后续专家意见、整改措施和闭环说明。", size=16)
    set_shape_text(slide.shapes[5], "[专家意见与整改对照表占位]", size=16, bold=True, color=MUTED, align=PP_ALIGN.CENTER)

    prs.save(str(output_path))


if __name__ == "__main__":
    main()
