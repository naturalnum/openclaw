#!/usr/bin/env python3
"""Generate a preview PPT demonstrating chart-heavy project-report slides."""

from __future__ import annotations

import argparse
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


BLUE = RGBColor(20, 73, 158)
LIGHT_BLUE = RGBColor(232, 241, 252)
MID_BLUE = RGBColor(91, 140, 222)
ACCENT = RGBColor(0, 163, 224)
TEXT = RGBColor(41, 49, 66)
MUTED = RGBColor(108, 117, 125)
BG = RGBColor(246, 249, 253)
WHITE = RGBColor(255, 255, 255)
LINE = RGBColor(210, 221, 236)
GREEN = RGBColor(35, 160, 131)
ORANGE = RGBColor(243, 156, 18)
RED = RGBColor(217, 83, 79)


TASKS = [
    {
        "name": "课题1",
        "title": "人机协同数据知识增强架构研究",
        "budget": 130,
        "unit": "中国电力科学研究院有限公司",
        "goal": "形成多时间尺度检修计划与运行方式协调编排整体架构",
    },
    {
        "name": "课题2",
        "title": "配网检修风险主动发现与溯因技术研究",
        "budget": 130,
        "unit": "西安交通大学 / 中国电科院",
        "goal": "构建风险预测、主动发现和溯因分析方法",
    },
    {
        "name": "课题3",
        "title": "专家反馈趋优学习与知识更新技术研究",
        "budget": 130,
        "unit": "北京交通大学 / 中国电科院",
        "goal": "形成反馈驱动的智能体趋优训练与知识库更新机制",
    },
    {
        "name": "课题4",
        "title": "智能校核与编制优化模块开发应用",
        "budget": 135,
        "unit": "上海公司 / 中国电科院 / 信通产业集团",
        "goal": "完成模块研发、系统集成和地区试点应用验证",
    },
]


MILESTONES = [
    ("2024.01-03", "启动项目，形成实施方案"),
    ("2024.04-09", "开展联合建模、知识提取与风险预测研究"),
    ("2024.10-2025.03", "推进表征方法、趋优训练和知识库更新"),
    ("2025.04-09", "形成增强架构与模块原型，开展试点验证"),
    ("2025.10-12", "完成验收材料、专利论文与总结报告"),
]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Generate a preview project report PPT.")
    parser.add_argument("--output", required=True, help="Output PPTX file path.")
    return parser.parse_args()


def add_textbox(slide, left, top, width, height, text, size=18, bold=False,
               color=TEXT, align=PP_ALIGN.LEFT, fill=None, line=None, radius=False):
    shape_type = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    box = slide.shapes.add_shape(shape_type, left, top, width, height)
    box.fill.solid()
    box.fill.fore_color.rgb = fill or WHITE
    if line is None:
        box.line.color.rgb = fill or WHITE
    else:
        box.line.color.rgb = line
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_title_band(slide, chapter, title):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = BG
    band = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.62)
    )
    band.fill.solid()
    band.fill.fore_color.rgb = BLUE
    band.line.fill.background()
    chapter_box = add_textbox(
        slide, Inches(0.45), Inches(0.78), Inches(3.2), Inches(0.55),
        chapter, size=24, bold=True, color=BLUE, fill=BG, line=BG
    )
    title_box = add_textbox(
        slide, Inches(0.45), Inches(1.28), Inches(6.5), Inches(0.55),
        title, size=16, color=MUTED, fill=BG, line=BG
    )
    for box in (chapter_box, title_box):
        box.line.fill.background()
    divider = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.45), Inches(1.95), Inches(12.4), Inches(0.03)
    )
    divider.fill.solid()
    divider.fill.fore_color.rgb = LINE
    divider.line.fill.background()


def add_upper_bullets(slide, bullets, emphasis=None):
    panel = add_textbox(
        slide, Inches(0.5), Inches(2.2), Inches(12.3), Inches(1.55),
        "", fill=WHITE, line=LINE, radius=True
    )
    tf = panel.text_frame
    tf.clear()
    tf.margin_left = Pt(18)
    tf.margin_right = Pt(18)
    tf.margin_top = Pt(10)
    tf.margin_bottom = Pt(10)
    for idx, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = f"• {bullet}"
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(18)
        p.font.color.rgb = TEXT
        p.space_after = Pt(6)
    if emphasis:
        badge = add_textbox(
            slide, Inches(10.35), Inches(2.45), Inches(2.0), Inches(0.75),
            emphasis, size=18, bold=True, color=WHITE, align=PP_ALIGN.CENTER,
            fill=ACCENT, line=ACCENT, radius=True
        )
        badge.line.fill.background()


def add_connector(slide, x1, y1, x2, y2, color=BLUE, width=2.5):
    connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    connector.line.color.rgb = color
    connector.line.width = Pt(width)
    return connector


def build_cover(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = BG
    top_band = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.0)
    )
    top_band.fill.solid()
    top_band.fill.fore_color.rgb = BLUE
    top_band.line.fill.background()

    add_textbox(
        slide, Inches(0.7), Inches(1.5), Inches(11.8), Inches(1.15),
        "基于人机协同的配网检修运行方式智能校核与编制优化关键技术研究",
        size=26, bold=True, color=BLUE, fill=BG, line=BG
    ).line.fill.background()
    add_textbox(
        slide, Inches(0.72), Inches(2.7), Inches(5.2), Inches(0.55),
        "项目汇报样稿（图表自动生成预览）",
        size=18, color=MUTED, fill=BG, line=BG
    ).line.fill.background()

    info = [
        ("承担单位", "国网上海市电力公司"),
        ("实施周期", "2024 年 1 月 - 2025 年 12 月"),
        ("项目经费", "525 万元"),
        ("课题数量", "4 个"),
    ]
    for idx, (label, value) in enumerate(info):
        left = Inches(0.8 + (idx % 2) * 6.0)
        top = Inches(4.2 + (idx // 2) * 1.1)
        add_textbox(
            slide, left, top, Inches(5.3), Inches(0.82),
            f"{label}\n{value}",
            size=18, bold=False, color=TEXT, fill=WHITE, line=LINE, radius=True
        )

    footer = add_textbox(
        slide, Inches(0.8), Inches(6.85), Inches(3.2), Inches(0.45),
        "Demo generated on 2026-04-07", size=12, color=MUTED, fill=BG, line=BG
    )
    footer.line.fill.background()


def build_goal_slide(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title_band(slide, "一、项目目标与总体思路", "用任务书内容自动生成上文下图型页面")
    add_upper_bullets(
        slide,
        [
            "面向配网多时间尺度检修运行方式智能校核与协同编制优化需求，构建数据知识融合的智能决策模型。",
            "解决人工提报、人工编排、人工协调中的数据复杂、约束缺失、过程低效等问题。",
            "支撑科学、合理、高效的检修计划与运行方式编制，降低设备停运和系统停电风险。",
        ],
        emphasis="目标导向页",
    )

    labels = [
        ("核心问题", "数据复杂\n约束缺失\n协同低效", BLUE),
        ("关键技术", "知识增强架构\n风险发现与溯因\n反馈趋优学习", MID_BLUE),
        ("模块研发", "智能校核\n编制优化\n系统集成", ACCENT),
        ("应用价值", "风险下降\n效率提升\n试点落地", GREEN),
    ]
    x_positions = [0.8, 3.55, 6.3, 9.05]
    for (title, body, color), x in zip(labels, x_positions):
        add_textbox(
            slide, Inches(x), Inches(4.25), Inches(2.35), Inches(1.75),
            f"{title}\n{body}", size=17, bold=False, color=WHITE,
            align=PP_ALIGN.CENTER, fill=color, line=color, radius=True
        )
    for idx in range(3):
        add_connector(
            slide,
            Inches(x_positions[idx] + 2.35), Inches(5.13),
            Inches(x_positions[idx + 1]), Inches(5.13),
            color=ACCENT if idx == 1 else BLUE,
        )


def build_task_slide(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title_band(slide, "二、四个课题分解", "下半区用课题卡片图展示研究内容和责任主体")
    add_upper_bullets(
        slide,
        [
            "任务书将项目拆分为 4 个课题，形成“架构研究 - 风险溯因 - 趋优学习 - 模块应用”的链式布局。",
            "每个课题均可自动抽取研究方向、承担单位、预算和预期目标，并映射为统一的图形卡片。",
        ],
        emphasis="课题拆解页",
    )

    positions = [
        (0.65, 4.15), (6.75, 4.15), (0.65, 5.55), (6.75, 5.55)
    ]
    for item, (x, y) in zip(TASKS, positions):
        add_textbox(
            slide, Inches(x), Inches(y), Inches(5.75), Inches(1.08),
            f"{item['name']}  {item['title']}\n{item['unit']}",
            size=15, bold=False, color=TEXT, fill=WHITE, line=LINE, radius=True
        )
        add_textbox(
            slide, Inches(x + 4.35), Inches(y + 0.14), Inches(1.1), Inches(0.55),
            f"{item['budget']}万", size=16, bold=True, color=WHITE, align=PP_ALIGN.CENTER,
            fill=BLUE, line=BLUE, radius=True
        )
        add_textbox(
            slide, Inches(x + 0.2), Inches(y + 0.53), Inches(4.0), Inches(0.42),
            item["goal"], size=11.5, color=MUTED, fill=WHITE, line=WHITE
        ).line.fill.background()


def build_budget_slide(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title_band(slide, "三、经费与分工", "下半区示例：原生形状绘制横向条形图")
    add_upper_bullets(
        slide,
        [
            "任务书中的承担单位和经费信息结构清晰，适合直接转为横向条形图或责任矩阵图。",
            "这类图建议优先用 PPT 原生形状生成，后续改色、调宽度、换单位都很方便。",
        ],
        emphasis="条形图页",
    )

    chart_left = Inches(1.0)
    chart_top = Inches(4.25)
    chart_width = 11.1
    max_budget = max(item["budget"] for item in TASKS)
    for idx, item in enumerate(TASKS):
        top = chart_top + Inches(idx * 0.56)
        add_textbox(
            slide, chart_left, top, Inches(3.6), Inches(0.36),
            f"{item['name']}  {item['title'][:18]}", size=13, color=TEXT, fill=BG, line=BG
        ).line.fill.background()
        bg_bar = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
            chart_left + Inches(3.8), top + Inches(0.02), Inches(5.7), Inches(0.28)
        )
        bg_bar.fill.solid()
        bg_bar.fill.fore_color.rgb = LIGHT_BLUE
        bg_bar.line.fill.background()
        bar_width = 5.7 * item["budget"] / max_budget
        fg_bar = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
            chart_left + Inches(3.8), top + Inches(0.02), Inches(bar_width), Inches(0.28)
        )
        fg_bar.fill.solid()
        fg_bar.fill.fore_color.rgb = BLUE if idx % 2 == 0 else ACCENT
        fg_bar.line.fill.background()
        add_textbox(
            slide, chart_left + Inches(9.7), top - Inches(0.02), Inches(1.0), Inches(0.36),
            f"{item['budget']} 万", size=13, bold=True, color=TEXT, fill=BG, line=BG
        ).line.fill.background()

    tip = add_textbox(
        slide, Inches(1.0), Inches(6.85), Inches(11.3), Inches(0.45),
        "说明：正式版可改为“承担单位经费分布”“课题研究内容经费分解”或“单位-任务-经费矩阵表”。",
        size=12, color=MUTED, fill=BG, line=BG
    )
    tip.line.fill.background()


def build_kpi_timeline_slide(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title_band(slide, "四、指标与进度", "下半区组合：KPI 卡片 + 项目时间轴")
    add_upper_bullets(
        slide,
        [
            "任务书中的硬指标非常适合做成数字卡片，突出验收关注点。",
            "进度计划可以进一步展开为按课题分层的甘特图，这是最典型的自动化场景之一。",
        ],
        emphasis="KPI/时间轴",
    )

    kpis = [
        ("风险研判类别", ">= 3 类", ORANGE),
        ("知识节点数", "> 10000", BLUE),
        ("安全校核通过率", "> 95%", GREEN),
        ("风险辨识准确率", "> 90%", ACCENT),
    ]
    for idx, (label, value, color) in enumerate(kpis):
        add_textbox(
            slide, Inches(0.8 + idx * 3.1), Inches(4.15), Inches(2.65), Inches(1.1),
            f"{label}\n{value}", size=18, bold=False, color=WHITE,
            align=PP_ALIGN.CENTER, fill=color, line=color, radius=True
        )

    start_x = 1.15
    y = 6.0
    axis = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(start_x), Inches(y), Inches(10.8), Inches(0.04)
    )
    axis.fill.solid()
    axis.fill.fore_color.rgb = LINE
    axis.line.fill.background()
    for idx, (period, text) in enumerate(MILESTONES):
        x = start_x + idx * 2.15
        marker = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.OVAL, Inches(x), Inches(y - 0.13), Inches(0.28), Inches(0.28)
        )
        marker.fill.solid()
        marker.fill.fore_color.rgb = BLUE if idx < 3 else ACCENT
        marker.line.fill.background()
        add_textbox(
            slide, Inches(x - 0.2), Inches(y - 0.72), Inches(1.2), Inches(0.32),
            period, size=11, bold=True, color=BLUE, align=PP_ALIGN.CENTER, fill=BG, line=BG
        ).line.fill.background()
        add_textbox(
            slide, Inches(x - 0.2), Inches(y + 0.2), Inches(1.8), Inches(0.7),
            text, size=10.5, color=TEXT, fill=BG, line=BG
        ).line.fill.background()


def build_system_slide(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title_band(slide, "五、系统应用示意", "下半区示例：分层架构图")
    add_upper_bullets(
        slide,
        [
            "课题 4 明确要求形成智能校核与编制优化模块，并与 PMS 等系统集成后开展地区试点应用。",
            "这类页面建议使用分层架构图，既能承接模板风格，也便于后续继续细化接口和组件。",
        ],
        emphasis="架构图页",
    )

    layers = [
        ("业务系统层", "PMS / 调控业务 / 检修计划报送", BLUE),
        ("智能模块层", "智能校核 / 风险辨识 / 预案生成 / 编制优化", MID_BLUE),
        ("知识与模型层", "先验知识 / 趋优训练 / 风险溯因 / 知识库更新", ACCENT),
        ("数据底座层", "设备台账 / 电网模型 / 运行数据 / 专家反馈", GREEN),
    ]
    for idx, (label, text, color) in enumerate(layers):
        add_textbox(
            slide, Inches(1.2), Inches(4.05 + idx * 0.72), Inches(10.9), Inches(0.52),
            f"{label}    {text}", size=16, color=WHITE, fill=color, line=color, radius=True
        )
    add_textbox(
        slide, Inches(11.35), Inches(4.05), Inches(1.0), Inches(2.7),
        "试点应用\n验证", size=18, bold=True, color=WHITE,
        align=PP_ALIGN.CENTER, fill=RED, line=RED, radius=True
    )


def main() -> None:
    args = parse_args()
    output_path = Path(args.output).expanduser().resolve()
    output_path.parent.mkdir(parents=True, exist_ok=True)

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    build_cover(prs)
    build_goal_slide(prs)
    build_task_slide(prs)
    build_budget_slide(prs)
    build_kpi_timeline_slide(prs)
    build_system_slide(prs)

    prs.save(str(output_path))


if __name__ == "__main__":
    main()
