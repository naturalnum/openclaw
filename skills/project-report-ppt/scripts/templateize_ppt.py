#!/usr/bin/env python3
"""Convert a finished report PPT into a reusable template."""

from __future__ import annotations

import argparse
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Pt


KEEP_EXACT = {
    "汇报内容",
    "汇 报 提 纲",
    "contents",
    "一、项目执行情况",
    "一、项目基本情况",
    "二、国内外研究进展",
    "三、主要研究内容及完成情况",
    "三、完成情况及主要创新成果",
    "四、总结和展望",
    "四、总结与展望",
    "五、专家意见回应",
    "项目概况",
    "项目组织管理",
    "关键事件节点",
    "变更事项",
    "经费使用情况",
    "整体研究内容",
    "立项背景",
    "项目展望",
    "中期督导情况",
}

KEEP_CONTAINS = [
    "一、项目基本情况",
    "二、国内外研究进展",
    "三、完成情况及主要创新成果",
    "三、主要研究内容及完成情况",
    "四、总结",
    "五、专家意见回应",
]

FIELD_PLACEHOLDERS = {
    "项目名称": "[项目名称]",
    "项目负责人：": "[项目负责人]",
    "牵头单位：": "[牵头单位]",
    "参与单位：": "[参与单位列表]",
    "项目负责人": "[项目负责人]",
    "项目牵头单位": "[项目牵头单位]",
    "出资单位": "[出资单位]",
    "项目经费": "[项目经费]",
    "起止时间": "[起止时间]",
    "课题名称": "[课题名称]",
    "研究单位": "[研究单位]",
    "单位": "[单位]",
    "总预算": "[总预算]",
    "累计支出": "[累计支出]",
    "预算执行率": "[预算执行率]",
}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Turn a finished PPT into a reusable template.")
    parser.add_argument("--input", required=True, help="Source PPTX.")
    parser.add_argument("--output", required=True, help="Output PPTX.")
    return parser.parse_args()


def normalize(text: str) -> str:
    return " ".join(text.split()).strip()


def should_keep(text: str) -> bool:
    cleaned = normalize(text)
    if not cleaned:
        return True
    if cleaned in KEEP_EXACT:
        return True
    return any(token in cleaned for token in KEEP_CONTAINS)


def classify_placeholder(shape, text: str, slide_index: int) -> str:
    cleaned = normalize(text)
    if slide_index == 0:
        if "年" in cleaned or "月" in cleaned:
            return "[汇报日期]"
        if "项目负责人" in cleaned:
            return "项目负责人： [请填写]"
        if "牵头单位" in cleaned:
            return "牵头单位： [请填写]"
        if "参与单位" in cleaned:
            return "参与单位： [请填写]"
        return "[项目名称]"

    if cleaned in FIELD_PLACEHOLDERS:
        return FIELD_PLACEHOLDERS[cleaned]

    left = shape.left
    top = shape.top
    width = shape.width
    height = shape.height

    if top < 1200000 and width > 3000000:
        return "[章节标题]"
    if top < 2600000 and height < 900000:
        return "[页面小标题]"
    if width > 5000000 and height > 1200000:
        return "[请替换本页正文内容]"
    if width < 1500000 and height < 700000:
        return "[标签]"
    if any(token in cleaned for token in ["课题", "指标", "论文", "专利", "系统", "原型", "示范"]):
        return "[请替换对应字段]"
    if any(ch.isdigit() for ch in cleaned):
        return "[数据占位]"
    return "[请替换内容]"


def set_text(shape, text: str) -> None:
    if not shape.has_text_frame:
        return
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER if len(text) <= 14 else PP_ALIGN.LEFT
    run = p.add_run()
    run.text = text
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(16 if len(text) <= 18 else 14)
    run.font.bold = False
    run.font.color.rgb = RGBColor(70, 80, 96)


def add_placeholder_box(slide, shape, label: str) -> None:
    box = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        shape.left,
        shape.top,
        shape.width,
        shape.height,
    )
    box.fill.solid()
    box.fill.fore_color.rgb = RGBColor(245, 247, 250)
    box.line.color.rgb = RGBColor(160, 170, 185)
    box.line.width = Pt(1.5)
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = label
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(16)
    run.font.bold = True
    run.font.color.rgb = RGBColor(120, 128, 140)


def remove_shape(shape) -> None:
    sp = shape._element
    sp.getparent().remove(sp)


def main() -> None:
    args = parse_args()
    input_path = Path(args.input).expanduser().resolve()
    output_path = Path(args.output).expanduser().resolve()
    output_path.parent.mkdir(parents=True, exist_ok=True)

    prs = Presentation(str(input_path))

    for slide_index, slide in enumerate(prs.slides):
        original_shapes = list(slide.shapes)
        for shape in reversed(original_shapes):
            if shape.shape_type in {MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.CHART}:
                label = "[图片占位]" if shape.shape_type == MSO_SHAPE_TYPE.PICTURE else "[图表占位]"
                add_placeholder_box(slide, shape, label)
                remove_shape(shape)
                continue

            if getattr(shape, "has_table", False):
                table = shape.table
                for row_idx, row in enumerate(table.rows):
                    for col_idx, cell in enumerate(row.cells):
                        if row_idx == 0:
                            cell.text = "[表头]"
                        elif col_idx == 0:
                            cell.text = "[项目]"
                        else:
                            cell.text = "[待填]"
                continue

            if not getattr(shape, "has_text_frame", False):
                continue

            text = shape.text or ""
            cleaned = normalize(text)
            if not cleaned:
                continue
            if should_keep(cleaned):
                continue

            placeholder = classify_placeholder(shape, cleaned, slide_index)
            set_text(shape, placeholder)

    prs.save(str(output_path))


if __name__ == "__main__":
    main()
